import os
import json
from datetime import datetime
import fitz  # PyMuPDF
import openpyxl
from dotenv import load_dotenv
load_dotenv()

PDFS_FOLDER = os.getenv("PDFS_FOLDER", "invoices")
OUTPUT_FOLDER = os.getenv("OUTPUT_FOLDER", "output")

# ─────────────────────────────────────────────
# WHITELIST — only these files can be read
# ─────────────────────────────────────────────
READABLE_FILES = ["main.py", "tools.py"]

# ─────────────────────────────────────────────
# TOOL DEFINITIONS (sent to Claude API)
# ─────────────────────────────────────────────
# HELP: https://platform.claude.com/docs/en/agents-and-tools/tool-use/define-tools
TOOLS = [
    {
        "name": "list_pdfs",
        "description": ("Lists all PDF files available in the configured PDFs folder. "
                        "Returns the folder path, total count, and filenames. "
                        "Use this tool when the user wants to know which PDFs are available, "
                        "before processing or selecting a specific file. Does not return file "
                        "contents or metadata beyond filenames."),
        "input_schema": {
            "type": "object",
            "properties": {},
            "required": []
        }
    },
    {
        "name": "extract_pdf_text",
        "description": (
            "Extracts the first 30 and last 30 lines of text from a PDF file. "
            "Use this to retrieve both header information (client name, address, date) "
            "and footer information (totals, tax breakdown, bank details, or payment terms) "
            "which may be located at the end of multi-page invoices."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "filename": {
                    "type": "string",
                    "description": "PDF filename (e.g. invoice_001.pdf)",
                }
            },
            "required": ["filename"],
        },
    }, 
    {
        "name": "create_excel",
        "description": (
            "Creates an Excel file (.xlsx) with the data extracted from the invoices. "
            "Call this tool at the end, after processing all PDFs."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "clients": {
                    "type": "array",
                    "description": "List of clients extracted from the invoices.",
                    "items": {
                        "type": "object",
                        "properties": {
                            "name":       {"type": "string", "description": "Client name"},
                            "address":    {"type": "string", "description": "Full address"},
                            "phone":      {"type": "string", "description": "Phone number"},
                            "amount_owed":{"type": "string", "description": "Total amount owed (e.g. 1250.00)"},
                        },
                        "required": ["name", "address", "phone", "amount_owed"],
                    },
                }
            },
            "required": ["clients"],
        },
    },  

    {
        "name": "read_project_code",
        "description": (
            "Reads the source code of the project files (main.py and tools.py). "
            "Use this tool ONLY when the user asks a question starting with /explain. "
            "This gives you full context of the codebase to answer technical questions accurately."
        ),
        "input_schema": {
            "type": "object",
            "properties": {},
            "required": [],
        },
    },              
]

# ─────────────────────────────────────────────
# TOOL EXECUTION
# ─────────────────────────────────────────────
def run_tool(tool_name: str, tool_input: dict) -> str:
    """Executes a tool by name and returns the result as a JSON string."""
    if tool_name == "list_pdfs":
        return _list_pdfs() 
    elif tool_name == "extract_pdf_text":
        return _extract_pdf_text(**tool_input)  
    elif tool_name == "create_excel":
        return _create_excel(tool_input["clients"])  
    elif tool_name == "read_project_code":
        return _read_project_code()                
    else:
        return json.dumps({"error": f"Unknown tool: {tool_name}"})
# ─────────────────────────────────────────────
# TOOLS
# ─────────────────────────────────────────────
def _list_pdfs() -> str:
    """Lists PDF files in the configured folder."""
    if not os.path.exists(PDFS_FOLDER):
        return json.dumps({
            "error": f"Folder '{PDFS_FOLDER}' not found.",
            "hint": f"Create a '{PDFS_FOLDER}' folder and place your PDFs inside."
        })

    files = [f for f in os.listdir(PDFS_FOLDER) if f.lower().endswith(".pdf")]

    if not files:
        return json.dumps({"error": f"No PDF files found in '{PDFS_FOLDER}'."})

    return json.dumps({
        "folder": PDFS_FOLDER,
        "total": len(files),
        "files": files
    })

def _extract_pdf_text(filename: str) -> str:
    """
    Extracts the first 30 and the last 30 non-empty lines from a PDF.
    Handles multi-page documents by aggregating lines before slicing.
    """
    path = os.path.join(PDFS_FOLDER, filename)

    if not os.path.exists(path):
        return json.dumps({"error": f"File '{filename}' not found."})

    try:
        all_lines = []
        
        # Use context manager to ensure the file is closed automatically
        with fitz.open(path) as doc:
            for page in doc:
                text = page.get_text()
                # Clean and collect non-empty lines
                page_lines = [line.strip() for line in text.splitlines() if line.strip()]
                all_lines.extend(page_lines)

        # Logic to handle short vs long documents
        if len(all_lines) <= 60:
            final_content = all_lines
        else:
            # Combine the top 30 and bottom 30 lines
            header = all_lines[:30]
            footer = all_lines[-30:]
            final_content = header + ["--- [Content Truncated] ---"] + footer

        return json.dumps({
            "file": filename,
            "text": "\n".join(final_content),
            "line_count": len(all_lines)
        })

    except Exception as e:
        return json.dumps({"error": f"Failed to read '{filename}': {str(e)}"})
    

def _create_excel(clients: list) -> str:
    """Creates the Excel file with client data."""
    os.makedirs(OUTPUT_FOLDER, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"clients_{timestamp}.xlsx"
    output_path = os.path.join(OUTPUT_FOLDER, filename)

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Clients"

        # Header row
        headers = ["Name", "Address", "Phone", "Amount Owed"]
        ws.append(headers)

        # Style the header
        from openpyxl.styles import Font, PatternFill, Alignment
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2E75B6")
            cell.alignment = Alignment(horizontal="center")

        # Data rows
        for client in clients:
            ws.append([
                client.get("name", ""),
                client.get("address", ""),
                client.get("phone", ""),
                client.get("amount_owed", ""),
            ])

        # Auto-fit column widths
        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

        wb.save(output_path)

        return json.dumps({
            "success": True,
            "file_created": filename,
            "path": output_path,
            "total_clients": len(clients)
        })

    except Exception as e:
        return json.dumps({"error": f"Failed to create Excel: {str(e)}"})
    
# ─────────────────────────────────────────────
# HELP
# ─────────────────────────────────────────────
def _read_project_code() -> str:
    """
    Reads the whitelisted project source files and returns their content.
    Only files listed in READABLE_FILES are accessible.
    """
    base_dir = os.path.dirname(os.path.abspath(__file__))
    result = {}
 
    for filename in READABLE_FILES:
        filepath = os.path.join(base_dir, filename)
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                result[filename] = f.read()
        except FileNotFoundError:
            result[filename] = f"[File not found: {filepath}]"
        except Exception as e:
            result[filename] = f"[Error reading file: {str(e)}]"
 
    return json.dumps({
        "files_read": list(result.keys()),
        "source_code": result
    })